#!/usr/bin/env python3
"""Generate the Consulting Services Agreement as a one-page, two-column .docx."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_columns(section, num=2, space_twips=360):
    """Set the number of columns on a section's sectPr."""
    sectPr = section._sectPr
    cols = sectPr.find(qn('w:cols'))
    if cols is None:
        cols = OxmlElement('w:cols')
        sectPr.append(cols)
    cols.set(qn('w:num'), str(num))
    cols.set(qn('w:space'), str(space_twips))


def no_space(p, before=0, after=2, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def heading(doc, text, size=8.5):
    p = doc.add_paragraph()
    no_space(p, before=3, after=1)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    r.font.name = 'Calibri'
    return p


def body(doc, text, size=8, bold=False):
    p = doc.add_paragraph()
    no_space(p, before=0, after=2)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.name = 'Calibri'
    return p


doc = Document()

# Tight margins to fit one page
sec = doc.sections[0]
sec.top_margin = Inches(0.5)
sec.bottom_margin = Inches(0.5)
sec.left_margin = Inches(0.6)
sec.right_margin = Inches(0.6)

# ---- Full-width header (single column) ----
title = doc.add_paragraph()
no_space(title, before=0, after=1)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('CONSULTING SERVICES AGREEMENT')
r.bold = True
r.font.size = Pt(13)
r.font.name = 'Calibri'

intro = doc.add_paragraph()
no_space(intro, before=0, after=3)
intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = intro.add_run('This Consulting Services Agreement (the "Agreement") is entered into as of '
                  'May 29, 2026 by and between the parties below.')
r.font.size = Pt(8)
r.font.name = 'Calibri'

# Parties block (full width, two short paragraphs)
for label, txt in [
    ('Client:  ', 'DIAMI USA CORPORATION, a corporation organized under the laws of the State of '
                  'California, with its principal office at 3435 Wilshire Blvd #2090, Los Angeles, CA 90010 '
                  '("Client"); and'),
    ('Consultant:  ', 'Jenna Armitage, doing business as ILYN Studio, an individual located at '
                      '2316 Hillcrest St, Suite 10, Orlando, FL 32803 ("Consultant").'),
]:
    p = doc.add_paragraph()
    no_space(p, before=0, after=1)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    rb = p.add_run(label)
    rb.bold = True
    rb.font.size = Pt(8)
    rb.font.name = 'Calibri'
    rt = p.add_run(txt)
    rt.font.size = Pt(8)
    rt.font.name = 'Calibri'

p = doc.add_paragraph()
no_space(p, before=0, after=2)
r = p.add_run('Client and Consultant are each a "Party" and together the "Parties."')
r.font.size = Pt(8)
r.font.name = 'Calibri'

# ---- Start two-column continuous section ----
new_sec = doc.add_section(WD_SECTION.CONTINUOUS)
new_sec.top_margin = Inches(0.5)
new_sec.bottom_margin = Inches(0.5)
new_sec.left_margin = Inches(0.6)
new_sec.right_margin = Inches(0.6)
set_columns(new_sec, num=2, space_twips=300)

clauses = [
    ("1. Engagement and Scope of Services",
     "Client engages Consultant as an independent consultant to provide professional consulting and "
     "advisory services in connection with Premiere Orlando 2026 (the \"Event\"), including but not "
     "limited to: on-site nail-art consulting, technical advisory, product and technique feedback, and "
     "related advisory support (the \"Services\"). Consultant shall perform the Services in a professional "
     "and workmanlike manner."),
    ("2. Term and Location",
     "Consultant shall provide the Services over two (2) days, on May 31, 2026 and June 1, 2026, at the "
     "Premiere Orlando 2026 venue in Orlando, Florida, and at such other times as mutually agreed."),
    ("3. Compensation",
     "As full consideration for the Services, Client shall pay Consultant a consulting fee of USD 2,000 "
     "(Two Thousand U.S. Dollars), payable within fifteen (15) days after completion of the Services and "
     "receipt of an invoice."),
    ("4. Independent Contractor",
     "Consultant is an independent contractor, not an employee, partner, or agent of Client. Consultant "
     "is solely responsible for all taxes (including self-employment taxes), insurance, and statutory "
     "obligations arising from the compensation paid hereunder. This Agreement does not create any "
     "employment relationship. Consultant shall provide a completed IRS Form W-9, and Client may report "
     "the compensation to tax authorities (e.g., IRS Form 1099-NEC) as required by law."),
    ("5. Confidentiality",
     "Consultant shall keep confidential all non-public information of Client obtained in connection with "
     "the Services, and shall not disclose or use it for any purpose other than performing the Services. "
     "This obligation survives termination of this Agreement."),
    ("6. Intellectual Property",
     "Any work product, designs, materials, or recommendations created by Consultant specifically for "
     "Client in the course of the Services shall be the sole property of Client as \"work made for hire,\" "
     "and Consultant hereby assigns all rights therein to Client. Consultant retains rights to her "
     "pre-existing know-how and general skills."),
    ("7. Independent Activities / Non-Exclusivity",
     "This Agreement is non-exclusive. Each Party may engage in other business activities, provided the "
     "confidentiality obligations are observed."),
    ("8. Termination",
     "Either Party may terminate this Agreement upon written notice if the other Party materially breaches "
     "this Agreement and fails to cure within five (5) days. Client shall pay for Services properly "
     "performed up to the termination date."),
    ("9. Governing Law and Dispute Resolution",
     "This Agreement shall be governed by and construed in accordance with the laws of the State of "
     "California, without regard to its conflict-of-laws principles. The Parties consent to the exclusive "
     "jurisdiction of the state and federal courts located in Los Angeles County, California."),
    ("10. Entire Agreement",
     "This Agreement constitutes the entire agreement between the Parties and supersedes all prior "
     "understandings. Any amendment must be in writing and signed by both Parties."),
]

for h, b in clauses:
    heading(doc, h)
    body(doc, b)

# ---- Signature block: back to single column ----
sig_sec = doc.add_section(WD_SECTION.CONTINUOUS)
sig_sec.top_margin = Inches(0.5)
sig_sec.bottom_margin = Inches(0.5)
sig_sec.left_margin = Inches(0.6)
sig_sec.right_margin = Inches(0.6)
set_columns(sig_sec, num=1)

wit = doc.add_paragraph()
no_space(wit, before=5, after=4)
r = wit.add_run('IN WITNESS WHEREOF, the Parties have executed this Agreement as of the date first written above.')
r.bold = True
r.font.size = Pt(8)
r.font.name = 'Calibri'

# Signature table
table = doc.add_table(rows=1, cols=2)
table.autofit = True
cell_l, cell_r = table.rows[0].cells

sig_lines_l = [
    ('Client — DIAMI USA CORPORATION', True),
    ('', False),
    ('Signature: ______________________', False),
    ('Name: Myung Hee Park', False),
    ('Title: CEO', False),
    ('Date: ______________________', False),
]
sig_lines_r = [
    ('Consultant — Jenna Armitage (ILYN Studio)', True),
    ('', False),
    ('Signature: ______________________', False),
    ('Name: Jenna Armitage', False),
    ('Date: ______________________', False),
]


def fill_cell(cell, lines):
    cell.paragraphs[0].text = ''
    first = True
    for txt, bold in lines:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        no_space(p, before=0, after=1)
        r = p.add_run(txt)
        r.bold = bold
        r.font.size = Pt(8)
        r.font.name = 'Calibri'


fill_cell(cell_l, sig_lines_l)
fill_cell(cell_r, sig_lines_r)

doc.save('/home/user/work/Consulting_Services_Agreement_DIAMI_Armitage.docx')
print('saved')
