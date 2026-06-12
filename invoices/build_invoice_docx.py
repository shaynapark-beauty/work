#!/usr/bin/env python3
"""Generate the DIAMI x Dahree seminar invoice as a Word (.docx) document."""

from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x6B, 0x72, 0x80)
ACCENT = RGBColor(0x11, 0x18, 0x27)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)


def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    m = OxmlElement('w:tcMar')
    for tag, val in (('top', top), ('bottom', bottom), ('start', left), ('end', right)):
        node = OxmlElement(f'w:{tag}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        m.append(node)
    tcPr.append(m)


def bottom_border(cell, color="CBD5E1", sz=8):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(sz))
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), color)
    borders.append(bottom)
    tcPr.append(borders)


def run(p, text, size=11, bold=False, color=INK, name='Calibri'):
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = name
    return r


def no_space(p, before=0, after=0, line=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        pf.line_spacing = line


doc = Document()

# Page margins
sec = doc.sections[0]
sec.top_margin = Inches(0.6)
sec.bottom_margin = Inches(0.6)
sec.left_margin = Inches(0.75)
sec.right_margin = Inches(0.75)

# ---------- Header (brand left / INVOICE right) ----------
header = doc.add_table(rows=1, cols=2)
header.autofit = False
header.columns[0].width = Inches(4.0)
header.columns[1].width = Inches(3.0)

# Left brand block
lc = header.cell(0, 0)
lc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
p = lc.paragraphs[0]
no_space(p, after=2)
run(p, "DIAMI USA CORPORATION", size=18, bold=True, color=ACCENT)
p2 = lc.add_paragraph()
no_space(p2, after=0, line=1.25)
run(p2, "3435 Wilshire Blvd #2090\nLos Angeles, CA 90010\nUnited States", size=9.5, color=MUTED)

# Right invoice block
rc = header.cell(0, 1)
rc.vertical_alignment = WD_ALIGN_VERTICAL.TOP
p = rc.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
no_space(p, after=4)
run(p, "INVOICE", size=26, bold=True, color=ACCENT)
for label, val in (("Invoice No.", "DIAMI-2026-0530"),
                   ("Invoice Date", "June 12, 2026"),
                   ("Due Date", "Upon Receipt")):
    pp = rc.add_paragraph()
    pp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    no_space(pp, after=1, line=1.1)
    run(pp, f"{label}  ", size=9.5, color=MUTED)
    run(pp, val, size=9.5, bold=True, color=INK)

# Accent rule under header
rule = doc.add_paragraph()
no_space(rule, before=4, after=10)
pPr = rule._p.get_or_add_pPr()
pbdr = OxmlElement('w:pBdr')
bottom = OxmlElement('w:bottom')
bottom.set(qn('w:val'), 'single')
bottom.set(qn('w:sz'), '18')
bottom.set(qn('w:space'), '1')
bottom.set(qn('w:color'), '111827')
pbdr.append(bottom)
pPr.append(pbdr)

# ---------- Parties ----------
parties = doc.add_table(rows=1, cols=2)
parties.columns[0].width = Inches(3.5)
parties.columns[1].width = Inches(3.5)


def party(cell, label, who, addr):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    p = cell.paragraphs[0]
    no_space(p, after=3)
    run(p, label, size=8.5, bold=True, color=MUTED)
    p2 = cell.add_paragraph()
    no_space(p2, after=2)
    run(p2, who, size=11.5, bold=True, color=INK)
    p3 = cell.add_paragraph()
    no_space(p3, after=0, line=1.2)
    run(p3, addr, size=10, color=RGBColor(0x37, 0x41, 0x51))


party(parties.cell(0, 0), "BILL TO", "Dahree",
      "1257 Lakeside Drive\nSunnyvale, CA 94085\nUnited States")
party(parties.cell(0, 1), "FROM", "DIAMI USA Corporation",
      "3435 Wilshire Blvd #2090\nLos Angeles, CA 90010\nUnited States")

doc.add_paragraph().paragraph_format.space_after = Pt(6)

# ---------- Line items table ----------
items = doc.add_table(rows=2, cols=4)
items.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = [Inches(3.7), Inches(0.7), Inches(1.3), Inches(1.3)]
for i, w in enumerate(widths):
    items.columns[i].width = w

# header row
headers = ["DESCRIPTION", "QTY", "UNIT PRICE", "AMOUNT"]
aligns = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT,
          WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
for i, (h, a) in enumerate(zip(headers, aligns)):
    c = items.rows[0].cells[i]
    c.width = widths[i]
    set_cell_bg(c, "111827")
    set_cell_margins(c)
    p = c.paragraphs[0]
    p.alignment = a
    no_space(p)
    run(p, h, size=9, bold=True, color=WHITE)

# data row
row = items.rows[1]
for i in range(4):
    row.cells[i].width = widths[i]
    set_cell_margins(row.cells[i], top=120, bottom=120)
    row.cells[i].vertical_alignment = WD_ALIGN_VERTICAL.TOP
    bottom_border(row.cells[i], color="E5E7EB")

# description cell
dc = row.cells[0]
p = dc.paragraphs[0]
no_space(p, after=3)
run(p, "DIAMI x Dahree Nail Product Seminar — Seminar Fee", size=10.5, bold=True, color=INK)
p2 = dc.add_paragraph()
no_space(p2, after=0, line=1.2)
run(p2, "Venue: 8738 International Drive, Orlando, Florida 32819\nSeminar Date: May 30, 2026",
    size=9, color=MUTED)

for i, val in ((1, "1"), (2, "$2,281.76"), (3, "$2,281.76")):
    c = row.cells[i]
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    no_space(p)
    run(p, val, size=10.5, color=INK)

doc.add_paragraph().paragraph_format.space_after = Pt(2)

# ---------- Totals (right aligned) ----------
totals = doc.add_table(rows=3, cols=2)
totals.alignment = WD_TABLE_ALIGNMENT.RIGHT
totals.columns[0].width = Inches(1.6)
totals.columns[1].width = Inches(1.5)

rows_data = [("Subtotal", "$2,281.76", False),
             ("Tax", "$0.00", False),
             ("Total Due (USD)", "$2,281.76", True)]
for ri, (label, val, grand) in enumerate(rows_data):
    lc = totals.rows[ri].cells[0]
    vc = totals.rows[ri].cells[1]
    lc.width = Inches(1.6)
    vc.width = Inches(1.5)
    set_cell_margins(lc, top=60, bottom=60)
    set_cell_margins(vc, top=60, bottom=60)
    lp = lc.paragraphs[0]
    vp = vc.paragraphs[0]
    lp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    vp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    no_space(lp); no_space(vp)
    if grand:
        run(lp, label, size=13, bold=True, color=ACCENT)
        run(vp, val, size=13, bold=True, color=ACCENT)
        for c in (lc, vc):
            tcPr = c._tc.get_or_add_tcPr()
            borders = OxmlElement('w:tcBorders')
            top = OxmlElement('w:top')
            top.set(qn('w:val'), 'single')
            top.set(qn('w:sz'), '16')
            top.set(qn('w:space'), '0')
            top.set(qn('w:color'), '111827')
            borders.append(top)
            tcPr.append(borders)
    else:
        run(lp, label, size=10.5, color=MUTED)
        run(vp, val, size=10.5, color=INK)

doc.add_paragraph().paragraph_format.space_after = Pt(10)

# ---------- Payment instructions box ----------
paybox = doc.add_table(rows=1, cols=1)
pcell = paybox.cell(0, 0)
set_cell_bg(pcell, "F9FAFB")
set_cell_margins(pcell, top=160, bottom=160, left=200, right=200)
# border around box
tcPr = pcell._tc.get_or_add_tcPr()
borders = OxmlElement('w:tcBorders')
for edge in ('top', 'bottom', 'start', 'end'):
    e = OxmlElement(f'w:{edge}')
    e.set(qn('w:val'), 'single')
    e.set(qn('w:sz'), '6')
    e.set(qn('w:space'), '0')
    e.set(qn('w:color'), 'E5E7EB')
    borders.append(e)
tcPr.append(borders)

p = pcell.paragraphs[0]
no_space(p, after=8)
run(p, "PAYMENT INSTRUCTIONS — WIRE / ACH TRANSFER", size=9, bold=True, color=MUTED)

fields = ["Bank Name", "Account Name", "Account Number", "Routing Number",
          "SWIFT / BIC", "Bank Address", "Reference / Memo", "Other"]
for f in fields:
    fp = pcell.add_paragraph()
    no_space(fp, after=4, line=1.0)
    run(fp, f"{f}:  ", size=10, bold=True, color=RGBColor(0x37, 0x41, 0x51))
    # blank underlined space to fill in
    run(fp, " " * 60, size=10, color=INK).font.underline = True

doc.add_paragraph().paragraph_format.space_after = Pt(8)

# ---------- Notes ----------
notes = doc.add_paragraph()
no_space(notes, after=0, line=1.3)
run(notes, "Please reference Invoice No. DIAMI-2026-0530 with your payment. "
           "All amounts are stated in U.S. Dollars (USD). For any questions regarding "
           "this invoice, please contact DIAMI USA Corporation.", size=9, color=MUTED)

# ---------- Signature ----------
doc.add_paragraph().paragraph_format.space_after = Pt(24)
sig = doc.add_table(rows=1, cols=1)
sig.alignment = WD_TABLE_ALIGNMENT.RIGHT
scell = sig.cell(0, 0)
scell.width = Inches(2.6)
bottom_border(scell, color="9CA3AF", sz=8)
sp = scell.paragraphs[0]
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
no_space(sp)
run(sp, " ", size=10)
cap = doc.add_paragraph()
cap.alignment = WD_ALIGN_PARAGRAPH.RIGHT
no_space(cap, before=2)
run(cap, "Authorized Signature — DIAMI USA Corporation", size=9, color=MUTED)

# ---------- Thank you ----------
ty = doc.add_paragraph()
ty.alignment = WD_ALIGN_PARAGRAPH.CENTER
no_space(ty, before=20)
run(ty, "THANK YOU FOR YOUR BUSINESS", size=10, color=MUTED)

out = "/home/user/work/invoices/DIAMI-Dahree-Seminar-Invoice.docx"
doc.save(out)
print("Saved:", out)
